from __future__ import annotations

import io
import shutil
import tempfile
from unittest.mock import AsyncMock

import pytest
from docx import Document

from app.services import llm_cleaner as _lc
from app.services import pdf_extractor as _pe
from app.services.audiobook_service import AudiobookService
from app.services.audiobook_store import AudiobookStore
from app.services.llm_providers import LLMCredential

CREDENTIAL = LLMCredential.gemini("test-key")


@pytest.fixture(autouse=True)
def isolated_audiobooks_dir(monkeypatch):
    tmp = tempfile.mkdtemp(prefix="ss_audiobooks_test_")

    class _PatchedSettings:
        @property
        def AUDIOBOOKS_DIR(self) -> str:
            return tmp

    monkeypatch.setattr("app.services.audiobook_store.settings", _PatchedSettings())
    AudiobookStore._reset_for_tests()
    yield tmp
    AudiobookStore._reset_for_tests()
    shutil.rmtree(tmp, ignore_errors=True)


def _make_book(file_ext: str, page_count: int, uses_llm: bool = True) -> str:
    title = f"Test.{file_ext}"
    bid = AudiobookStore.create_book(title)
    meta = AudiobookStore.initial_meta(
        bid, title, page_count, "kokoro", "af_bella", 1.0, {"cost_usd": 0.0}
    )
    meta["file_ext"] = file_ext
    meta["uses_gemini_cleanup"] = uses_llm
    AudiobookStore.write_meta(bid, meta)
    return bid


def _titles(bid: str) -> list[str]:
    return [s["title"] for s in AudiobookStore.read_meta(bid)["sections"]]


def test_fallback_sections_small_book_is_single_section():
    assert AudiobookService._fallback_sections(10, "My Book") == [
        {"title": "My Book", "start_page": 1, "end_page": 10}
    ]


def test_fallback_sections_untitled_book_uses_default_title():
    assert AudiobookService._fallback_sections(5, None)[0]["title"] == "Audiobook"


def test_fallback_sections_threshold_is_inclusive():
    size = AudiobookService._FALLBACK_SECTION_PAGES

    assert len(AudiobookService._fallback_sections(size, "T")) == 1
    assert len(AudiobookService._fallback_sections(size + 1, "T")) == 2


def test_fallback_sections_large_book_chunks_into_contiguous_parts():
    sections = AudiobookService._fallback_sections(40, "Long Book")

    assert sections[0] == {"title": "Part 1", "start_page": 1, "end_page": 15}
    assert [s["title"] for s in sections] == ["Part 1", "Part 2", "Part 3"]
    assert sections[-1]["end_page"] == 40
    for previous, current in zip(sections, sections[1:], strict=False):
        assert current["start_page"] == previous["end_page"] + 1


@pytest.mark.asyncio
async def test_large_book_with_failed_detection_falls_back_to_parts(monkeypatch):
    bid = _make_book("txt", 40)
    detect = AsyncMock(side_effect=RuntimeError("simulated failure"))
    monkeypatch.setattr(_lc.LLMCleaner, "detect_sections", detect)

    await AudiobookService._phase_section(
        bid, CREDENTIAL, AudiobookStore.read_meta(bid)
    )

    persisted = AudiobookStore.read_meta(bid)
    assert persisted["sections_source"] == "fallback"
    assert len(persisted["sections"]) == 3
    assert persisted["sections"][0]["start_page"] == 1
    assert persisted["sections"][-1]["end_page"] == 40


@pytest.mark.asyncio
async def test_small_book_with_failed_detection_stays_single_section(monkeypatch):
    bid = _make_book("txt", 3)
    monkeypatch.setattr(
        _lc.LLMCleaner, "detect_sections", AsyncMock(side_effect=RuntimeError("x"))
    )

    await AudiobookService._phase_section(
        bid, CREDENTIAL, AudiobookStore.read_meta(bid)
    )

    sections = AudiobookStore.read_meta(bid)["sections"]
    assert len(sections) == 1
    assert sections[0]["end_page"] == 3


@pytest.mark.asyncio
async def test_large_local_first_text_book_without_structure_gets_parts():
    bid = _make_book("txt", 20, uses_llm=False)

    await AudiobookService._phase_section(bid, None, AudiobookStore.read_meta(bid))

    assert _titles(bid) == ["Part 1", "Part 2"]


@pytest.mark.asyncio
async def test_pdf_outline_short_circuits_the_llm(monkeypatch):
    bid = _make_book("pdf", 5)
    outline = [
        {"title": "Real Chapter One", "start_page": 1, "level": 0},
        {"title": "Real Chapter Two", "start_page": 3, "level": 0},
    ]
    monkeypatch.setattr(
        _pe.PDFExtractor, "read_outline", classmethod(lambda cls, path: outline)
    )
    detect = AsyncMock()
    monkeypatch.setattr(_lc.LLMCleaner, "detect_sections", detect)

    await AudiobookService._phase_section(
        bid, CREDENTIAL, AudiobookStore.read_meta(bid)
    )

    detect.assert_not_called()
    sections = AudiobookStore.read_meta(bid)["sections"]
    assert [(s["title"], s["start_page"], s["end_page"]) for s in sections] == [
        ("Real Chapter One", 1, 2),
        ("Real Chapter Two", 3, 5),
    ]


@pytest.mark.asyncio
async def test_outline_entries_outside_the_page_range_are_dropped(monkeypatch):
    bid = _make_book("pdf", 5)
    outline = [
        {"title": "Real Chapter", "start_page": 2, "level": 0},
        {"title": "Out Of Range", "start_page": 99, "level": 0},
        {"title": "Before Start", "start_page": 0, "level": 0},
    ]
    monkeypatch.setattr(
        _pe.PDFExtractor, "read_outline", classmethod(lambda cls, path: outline)
    )

    await AudiobookService._phase_section(
        bid, CREDENTIAL, AudiobookStore.read_meta(bid)
    )

    titles = _titles(bid)
    assert "Real Chapter" in titles
    assert "Out Of Range" not in titles
    assert "Before Start" not in titles


@pytest.mark.asyncio
async def test_docx_headings_short_circuit_the_llm(monkeypatch):
    bid = _make_book("docx", 1)
    doc = Document()
    doc.add_heading("Real Chapter One", level=1)
    doc.add_paragraph("Body text.")
    doc.add_heading("Real Chapter Two", level=1)
    doc.add_paragraph("More body text.")
    buffer = io.BytesIO()
    doc.save(buffer)
    AudiobookStore.save_source(bid, buffer.getvalue(), "docx")
    detect = AsyncMock()
    monkeypatch.setattr(_lc.LLMCleaner, "detect_sections", detect)

    await AudiobookService._phase_section(
        bid, CREDENTIAL, AudiobookStore.read_meta(bid)
    )

    detect.assert_not_called()
    persisted = AudiobookStore.read_meta(bid)
    assert persisted["sections_source"] == "detected"
    assert _titles(bid) == ["Real Chapter One", "Real Chapter Two"]


@pytest.mark.asyncio
async def test_docx_headings_are_used_for_local_first_books_too():
    bid = _make_book("docx", 1, uses_llm=False)
    doc = Document()
    doc.add_heading("Only Chapter", level=1)
    doc.add_paragraph("Body text.")
    buffer = io.BytesIO()
    doc.save(buffer)
    AudiobookStore.save_source(bid, buffer.getvalue(), "docx")

    await AudiobookService._phase_section(bid, None, AudiobookStore.read_meta(bid))

    assert _titles(bid) == ["Only Chapter"]


@pytest.mark.asyncio
async def test_markdown_headers_short_circuit_the_llm(monkeypatch):
    bid = _make_book("md", 1)
    AudiobookStore.save_source(
        bid,
        b"# Real Chapter One\n\nBody text.\n\n# Real Chapter Two\n\nMore text.\n",
        "md",
    )
    detect = AsyncMock()
    monkeypatch.setattr(_lc.LLMCleaner, "detect_sections", detect)

    await AudiobookService._phase_section(
        bid, CREDENTIAL, AudiobookStore.read_meta(bid)
    )

    detect.assert_not_called()
    assert _titles(bid) == ["Real Chapter One", "Real Chapter Two"]


@pytest.mark.asyncio
async def test_local_first_markdown_keeps_every_heading_level():
    bid = _make_book("md", 1, uses_llm=False)
    AudiobookStore.save_source(bid, b"Intro.\n\n### Deep Heading\n\nBody.\n", "md")

    await AudiobookService._phase_section(bid, None, AudiobookStore.read_meta(bid))

    assert "Deep Heading" in _titles(bid)


@pytest.mark.asyncio
async def test_markdown_without_h1_or_h2_falls_through_to_the_llm(monkeypatch):
    bid = _make_book("md", 1)
    AudiobookStore.save_source(bid, b"Intro.\n\n### Deep Heading\n\nBody.\n", "md")
    detect = AsyncMock(
        return_value=[{"title": "LLM Chapter", "start_page": 1, "end_page": 1}]
    )
    monkeypatch.setattr(_lc.LLMCleaner, "detect_sections", detect)

    await AudiobookService._phase_section(
        bid, CREDENTIAL, AudiobookStore.read_meta(bid)
    )

    detect.assert_called_once()
    assert _titles(bid) == ["LLM Chapter"]


@pytest.mark.asyncio
async def test_existing_detected_sections_are_not_recomputed(monkeypatch):
    bid = _make_book("pdf", 5)
    meta = AudiobookStore.read_meta(bid)
    meta["sections"] = [
        {"title": "Kept", "start_page": 1, "end_page": 5, "start_time": 0.0}
    ]
    meta["sections_source"] = "detected"
    AudiobookStore.write_meta(bid, meta)
    read_outline = AsyncMock()
    monkeypatch.setattr(
        _pe.PDFExtractor, "read_outline", classmethod(lambda cls, path: read_outline())
    )

    await AudiobookService._phase_section(
        bid, CREDENTIAL, AudiobookStore.read_meta(bid)
    )

    read_outline.assert_not_called()
    assert _titles(bid) == ["Kept"]
